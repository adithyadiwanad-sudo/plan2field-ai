import csv,io,zipfile
def parse(data,mapping):
    if not mapping or 'text' not in mapping:raise ValueError('Explicit text-column mapping is required')
    records=[]
    if data[:2]==b'PK':
        archive=zipfile.ZipFile(io.BytesIO(data))
        if sum(x.file_size for x in archive.infolist())>30_000_000 or len(archive.infolist())>500:raise ValueError('Workbook expansion limit exceeded')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        for sheet in book:
            rows=sheet.iter_rows(values_only=True);headers=next(rows,())
            if mapping['text'] not in headers:raise ValueError('Mapped text column not found')
            col=headers.index(mapping['text'])
            for index,row in enumerate(rows,2):
                if len(records)>=5000:raise ValueError('Report row limit exceeded')
                value=row[col]
                if isinstance(value,str) and value.startswith('='):raise ValueError('Formula evidence is not accepted')
                if value:records.append({'text':str(value),'sheet':sheet.title,'row':index})
    else:
        for index,row in enumerate(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))),2):
            if index>5001:raise ValueError('Report row limit exceeded')
            if mapping['text'] not in row:raise ValueError('Mapped text column not found')
            records.append({'text':row[mapping['text']],'sheet':'CSV','row':index})
    return records
